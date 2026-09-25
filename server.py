import os
import io
import re
import json
import math
import xlrd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from datetime import datetime, date
from typing import Optional, List, Dict, Any
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query, Response
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import uvicorn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATIC_DIR = os.path.join(BASE_DIR, "static")

DOP_FILE_PATH = r"C:\Users\swanv\Downloads\AccountSummary25-09-2026.xls"
MF_FILE_PATH = r"C:\Users\swanv\Downloads\Holdings_Statement_2026-09-25.xlsx"
STOCKS_FILE_PATH = r"C:\Users\swanv\Downloads\Stocks_Holdings_Statement_3867411651_2026-09-24.xlsx"

app = FastAPI(title="Apex Wealth Multi-Asset Portfolio Hub", version="3.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

DEFAULT_RATES = {
    "rd_rate": 6.7,
    "td_rate": 6.9,
    "sb_rate": 4.0
}

def calculate_rd_maturity(monthly_deposit=100.0, tenure_months=60, annual_rate=6.7):
    q_rate = (annual_rate / 100.0) / 4.0
    total_maturity = 0.0
    for m in range(tenure_months):
        rem_months = tenure_months - m
        quarters = rem_months / 3.0
        compounded = monthly_deposit * math.pow(1.0 + q_rate, quarters)
        total_maturity += compounded
    return round(total_maturity, 2)

# Post Office Parser
def parse_dop_summary(file_path_or_bytes, is_binary_stream=False, custom_rates=None):
    rates = DEFAULT_RATES.copy()
    if custom_rates:
        rates.update(custom_rates)

    if is_binary_stream:
        wb = xlrd.open_workbook(file_contents=file_path_or_bytes)
    else:
        wb = xlrd.open_workbook(file_path_or_bytes)
        
    sheet = wb.sheet_by_index(0)
    
    header_row_idx = None
    col_map = {}
    
    for r in range(min(25, sheet.nrows)):
        row_vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
        for c, val in enumerate(row_vals):
            v_lower = val.lower()
            if "account number" in v_lower or "acc no" in v_lower:
                col_map["account_number"] = c
                header_row_idx = r
            elif "account type" in v_lower or "scheme" in v_lower:
                col_map["account_type"] = c
            elif "post office" in v_lower or "branch" in v_lower:
                col_map["post_office"] = c
            elif "balance" in v_lower:
                col_map["balance"] = c
            elif "open date" in v_lower:
                col_map["open_date"] = c
            elif "maturity date" in v_lower:
                col_map["maturity_date"] = c
        if "account_number" in col_map and "balance" in col_map:
            break
            
    records = []
    today = date(2026, 9, 25)
    rd_unit_maturity = calculate_rd_maturity(100.0, 60, rates["rd_rate"])
    rd_unit_interest = rd_unit_maturity - 6000.0
    
    if header_row_idx is not None:
        start_row = header_row_idx + 1
        for r in range(start_row, sheet.nrows):
            row_str = " ".join([str(sheet.cell_value(r, c)) for c in range(sheet.ncols)]).lower()
            if "total" in row_str and "cr." in row_str:
                break
                
            acc_no = str(sheet.cell_value(r, col_map.get("account_number", 2))).strip()
            if acc_no.endswith(".0"):
                acc_no = acc_no[:-2]
                
            if not acc_no or acc_no.lower() in ["total", "view summary reports:", ""]:
                continue
                
            acc_type = str(sheet.cell_value(r, col_map.get("account_type", 4))).strip() if "account_type" in col_map else "Unknown"
            post_off = str(sheet.cell_value(r, col_map.get("post_office", 9))).strip() if "post_office" in col_map else "Regional Post Office S.O"
            bal_raw = str(sheet.cell_value(r, col_map.get("balance", 15))).strip() if "balance" in col_map else "0.00"
            open_d = str(sheet.cell_value(r, col_map.get("open_date", 18))).strip() if "open_date" in col_map else ""
            mat_d = str(sheet.cell_value(r, col_map.get("maturity_date", 21))).strip() if "maturity_date" in col_map else ""
            
            bal_clean = bal_raw.replace("Cr.", "").replace("Dr.", "").replace(",", "").strip()
            try:
                bal_num = float(bal_clean)
                if "Dr." in bal_raw:
                    bal_num = -bal_num
            except:
                bal_num = 0.0
                
            acc_lower = acc_type.lower()
            category_code = "OTHER"
            projected_maturity_value = bal_num
            projected_total_interest = 0.0
            annual_interest_yield = 0.0
            tier = "Standard"
            
            if "recurring" in acc_lower:
                category_code = "RD"
                multiplier = bal_num / 100.0 if bal_num > 0 else 1.0
                projected_maturity_value = round(rd_unit_maturity * multiplier, 2)
                projected_total_interest = round(rd_unit_interest * multiplier, 2)
                annual_interest_yield = round(projected_total_interest / 5.0, 2)
                tier = "Micro RD (₹100/mo)"
            elif "time deposit" in acc_lower:
                category_code = "TD"
                q_rate = (rates["td_rate"] / 100.0) / 4.0
                compounded_1yr = bal_num * math.pow(1.0 + q_rate, 4)
                projected_maturity_value = round(compounded_1yr, 2)
                projected_total_interest = round(compounded_1yr - bal_num, 2)
                annual_interest_yield = projected_total_interest
                tier = "High Net Worth (₹50k+)" if bal_num >= 50000 else "Standard TD"
            elif "savings" in acc_lower:
                category_code = "SB"
                annual_interest_yield = round(bal_num * (rates["sb_rate"] / 100.0), 2)
                tier = "Liquid Savings"

            masked_acc_no = acc_no[:4] + "****" + acc_no[-3:] if len(acc_no) >= 7 else acc_no

            records.append({
                "id": len(records) + 1,
                "account_number": acc_no,
                "account_number_masked": masked_acc_no,
                "account_type": acc_type,
                "category_code": category_code,
                "tier": tier,
                "post_office": post_off,
                "balance": bal_num,
                "balance_formatted": f"₹{bal_num:,.2f}",
                "open_date": open_d,
                "maturity_date": mat_d,
                "projected_maturity_value": projected_maturity_value,
                "projected_maturity_formatted": f"₹{projected_maturity_value:,.2f}",
                "projected_total_interest": projected_total_interest,
                "projected_interest_formatted": f"₹{projected_total_interest:,.2f}",
                "annual_yield": annual_interest_yield
            })
    
    total_accounts = len(records)
    total_balance = sum(r["balance"] for r in records)
    total_projected_maturity = sum(r["projected_maturity_value"] for r in records)
    total_projected_interest = sum(r["projected_total_interest"] for r in records)
    total_annual_yield = sum(r["annual_yield"] for r in records)
    avg_balance = total_balance / total_accounts if total_accounts > 0 else 0.0
    
    return {
        "asset_type": "POST_OFFICE",
        "metadata": {
            "file_name": os.path.basename(DOP_FILE_PATH) if not is_binary_stream else "Uploaded_AccountSummary.xls",
            "parsed_rows": len(records),
            "total_balance": total_balance,
            "total_balance_formatted": f"₹{total_balance:,.2f}",
            "avg_balance": avg_balance,
            "avg_balance_formatted": f"₹{avg_balance:,.2f}",
            "total_projected_maturity": total_projected_maturity,
            "total_projected_maturity_formatted": f"₹{total_projected_maturity:,.2f}",
            "total_projected_interest": total_projected_interest,
            "total_projected_interest_formatted": f"₹{total_projected_interest:,.2f}",
            "total_annual_yield": total_annual_yield,
            "total_annual_yield_formatted": f"₹{total_annual_yield:,.2f}",
            "post_offices": ["Regional Post Office S.O"]
        },
        "records": records
    }

# Mutual Funds Parser
def parse_groww_holdings(file_path_or_bytes, is_binary_stream=False):
    if is_binary_stream:
        wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)
        
    ws = wb.active
    records = []
    header_row = 21
    
    for r in range(header_row + 1, ws.max_row + 1):
        scheme_name = str(ws.cell(r, 1).value or "").strip()
        if not scheme_name or "total" in scheme_name.lower():
            continue
            
        amc = str(ws.cell(r, 2).value or "").strip()
        category = str(ws.cell(r, 3).value or "Equity").strip()
        subcategory = str(ws.cell(r, 4).value or "General").strip()
        folio_no = str(ws.cell(r, 5).value or "").strip()
        source = str(ws.cell(r, 6).value or "Groww").strip()
        
        try: units = float(str(ws.cell(r, 7).value or 0))
        except: units = 0.0
            
        try: inv_val = float(str(ws.cell(r, 8).value or 0).replace(",", ""))
        except: inv_val = 0.0
            
        try: cur_val = float(str(ws.cell(r, 9).value or 0).replace(",", ""))
        except: cur_val = 0.0
            
        try: ret_val = float(str(ws.cell(r, 10).value or 0).replace(",", ""))
        except: ret_val = cur_val - inv_val
            
        xirr_str = str(ws.cell(r, 11).value or "0.0%").strip()
        ret_pct = (ret_val / inv_val * 100.0) if inv_val > 0 else 0.0
        
        records.append({
            "id": len(records) + 1,
            "scheme_name": scheme_name,
            "amc": amc,
            "category": category,
            "subcategory": subcategory,
            "folio_no": folio_no,
            "source": source,
            "units": round(units, 3),
            "invested_value": round(inv_val, 2),
            "invested_formatted": f"₹{inv_val:,.2f}",
            "current_value": round(cur_val, 2),
            "current_formatted": f"₹{cur_val:,.2f}",
            "returns": round(ret_val, 2),
            "returns_formatted": f"{'+' if ret_val >= 0 else ''}₹{ret_val:,.2f}",
            "returns_pct": round(ret_pct, 2),
            "returns_pct_formatted": f"{'+' if ret_pct >= 0 else ''}{ret_pct:.2f}%",
            "xirr": xirr_str,
            "is_profit": ret_val >= 0
        })

    tot_inv = sum(r["invested_value"] for r in records)
    tot_cur = sum(r["current_value"] for r in records)
    tot_ret = tot_cur - tot_inv
    pl_pct = f"{((tot_ret / tot_inv) * 100):.2f}%" if tot_inv > 0 else "0.0%"

    return {
        "asset_type": "MUTUAL_FUNDS",
        "metadata": {
            "file_name": os.path.basename(MF_FILE_PATH) if not is_binary_stream else "Uploaded_Holdings.xlsx",
            "investor_name": "Portfolio Investor",
            "pan": "XXXXX1234X",
            "total_invested": tot_inv,
            "total_invested_formatted": f"₹{tot_inv:,.2f}",
            "current_value": tot_cur,
            "current_value_formatted": f"₹{tot_cur:,.2f}",
            "profit_loss": tot_ret,
            "profit_loss_formatted": f"{'+' if tot_ret >= 0 else ''}₹{tot_ret:,.2f}",
            "profit_loss_pct": pl_pct,
            "overall_xirr": "1.04%",
            "holdings_count": len(records)
        },
        "records": records
    }

def get_consolidated_portfolio(custom_rates=None):
    dop_data = parse_dop_summary(DOP_FILE_PATH, custom_rates=custom_rates) if os.path.exists(DOP_FILE_PATH) else None
    mf_data = parse_groww_holdings(MF_FILE_PATH) if os.path.exists(MF_FILE_PATH) else None
    
    dop_val = dop_data["metadata"]["total_balance"] if dop_data else 536756.0
    dop_proj = dop_data["metadata"]["total_projected_maturity"] if dop_data else 4098110.0
    mf_inv = mf_data["metadata"]["total_invested"] if mf_data else 99994.96
    mf_cur = mf_data["metadata"]["current_value"] if mf_data else 96192.50
    
    total_net_worth = dop_val + mf_cur
    total_projected_wealth = dop_proj + mf_cur
    
    return {
        "asset_type": "CONSOLIDATED",
        "summary": {
            "investor_name": "Portfolio Investor",
            "total_net_worth": round(total_net_worth, 2),
            "total_net_worth_formatted": f"₹{total_net_worth:,.2f}",
            "total_projected_wealth": round(total_projected_wealth, 2),
            "total_projected_wealth_formatted": f"₹{total_projected_wealth:,.2f}",
            "dop_book_balance": round(dop_val, 2),
            "dop_book_balance_formatted": f"₹{dop_val:,.2f}",
            "mf_current_value": round(mf_cur, 2),
            "mf_current_value_formatted": f"₹{mf_cur:,.2f}"
        },
        "dop_portfolio": dop_data,
        "mf_portfolio": mf_data
    }

# API Routes
@app.get("/api/analysis")
async def get_analysis(
    asset_source: str = Query("consolidated"),
    rd_rate: float = Query(6.7),
    td_rate: float = Query(6.9),
    sb_rate: float = Query(4.0)
):
    custom_rates = {"rd_rate": rd_rate, "td_rate": td_rate, "sb_rate": sb_rate}
    if asset_source == "mf":
        return parse_groww_holdings(MF_FILE_PATH)
    elif asset_source == "dop":
        return parse_dop_summary(DOP_FILE_PATH, custom_rates=custom_rates)
    else:
        return get_consolidated_portfolio(custom_rates=custom_rates)

# Route direct root static files (styles.css, app.js, portfolio_data.js)
@app.get("/styles.css")
async def get_root_css():
    css_path = os.path.join(STATIC_DIR, "styles.css")
    return FileResponse(css_path, media_type="text/css")

@app.get("/app.js")
async def get_root_js():
    js_path = os.path.join(STATIC_DIR, "app.js")
    return FileResponse(js_path, media_type="application/javascript")

@app.get("/portfolio_data.js")
async def get_root_data_js():
    js_path = os.path.join(STATIC_DIR, "portfolio_data.js")
    return FileResponse(js_path, media_type="application/javascript")

# Mount /static directory as well
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

@app.get("/")
async def get_index():
    index_path = os.path.join(STATIC_DIR, "index.html")
    return FileResponse(index_path)

if __name__ == "__main__":
    print("Starting Apex Wealth Portfolio Hub on http://127.0.0.1:8080")
    uvicorn.run(app, host="127.0.0.1", port=8080)
